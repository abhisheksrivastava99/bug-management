import json
import os
import sqlite3
import subprocess
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]


class FixtureGeneratorTests(unittest.TestCase):
    def test_generator_is_deterministic_and_consistent(self) -> None:
        env = dict(os.environ)
        env["BM_FIXTURE_SEED"] = "17"
        subprocess.run([sys.executable, "scripts/generate_excel_fixtures.py"], cwd=ROOT, check=True, env=env)
        catalog_path = ROOT / "fixtures" / "catalog" / "scenario_catalog.json"
        first = json.loads(catalog_path.read_text(encoding="utf-8"))
        subprocess.run([sys.executable, "scripts/generate_excel_fixtures.py"], cwd=ROOT, check=True, env=env)
        second = json.loads(catalog_path.read_text(encoding="utf-8"))
        self.assertEqual(first, second)
        self.assertGreaterEqual(len(first["scenarios"]), 8)
        scenario_types = {scenario["scenario_type"] for scenario in first["scenarios"]}
        self.assertIn("duplicate_data", scenario_types)
        self.assertIn("invalid_type_format", scenario_types)
        self.assertIn("join_filter_miss", scenario_types)
        for scenario in first["scenarios"]:
            self.assertIn("root_cause_family", scenario)
            self.assertIn("mapping_rows", scenario)
            self.assertIn("sql_fixture", scenario)
            self.assertGreaterEqual(len(scenario["mapping_rows"]), 1)
            for row in scenario["mapping_rows"]:
                self.assertNotEqual(row["old_column_name"], row["new_column_name"])
                self.assertNotIn("mapping_status", row)

        workbook = load_workbook(ROOT / "fixtures" / "excel" / "column_mapping.xlsx", read_only=True)
        header = next(workbook.active.iter_rows(values_only=True))
        self.assertNotIn("mapping_status", header)
        self.assertEqual(
            header,
            (
                "old_system",
                "new_system",
                "old_table_name",
                "new_table_name",
                "old_column_name",
                "new_column_name",
                "data_type_old",
                "data_type_new",
            ),
        )

        sqlite_path = ROOT / "fixtures" / "sqlite" / "bug_management_demo.db"
        self.assertTrue(sqlite_path.exists())
        with sqlite3.connect(sqlite_path) as connection:
            cursor = connection.execute('SELECT COUNT(*) FROM "tentity_new"')
            self.assertGreater(cursor.fetchone()[0], 0)
            cursor = connection.execute('SELECT COUNT(*) FROM "tdupshipment_new"')
            self.assertGreater(cursor.fetchone()[0], 0)


if __name__ == "__main__":
    unittest.main()
